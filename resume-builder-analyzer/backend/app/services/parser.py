import re
import io
from typing import Dict, List, Optional, Tuple, Any
from dataclasses import dataclass, field
import pdfplumber
import docx

@dataclass
class ParsedDateRange:
    raw_text: str
    start_year: Optional[int] = None
    start_month: Optional[int] = None
    end_year: Optional[int] = None
    end_month: Optional[int] = None
    is_current: bool = False

@dataclass
class ParsedResume:
    raw_text: str
    pages: int = 1
    word_count: int = 0
    sections: Dict[str, str] = field(default_factory=dict)
    bullet_points: List[str] = field(default_factory=list)
    has_tables: bool = False
    has_images: bool = False
    is_scanned: bool = False
    is_multi_column: bool = False
    emails: List[str] = field(default_factory=list)
    phones: List[str] = field(default_factory=list)
    links: List[str] = field(default_factory=list)
    date_ranges: List[ParsedDateRange] = field(default_factory=list)
    layout_issues: List[str] = field(default_factory=list)

# Section heading keywords
SECTION_PATTERNS = {
    "summary": [
        r"^(professional\s+)?summary\b",
        r"^profile\b",
        r"^about(\s+me)?\b",
        r"^career\s+objective\b",
        r"^executive\s+summary\b"
    ],
    "experience": [
        r"^(work\s+|professional\s+)?experience\b",
        r"^employment(\s+history)?\b",
        r"^work\s+history\b",
        r"^career\s+history\b"
    ],
    "education": [
        r"^education(\s+and\s+training)?\b",
        r"^academic\s+background\b",
        r"^qualifications\b"
    ],
    "skills": [
        r"^(technical\s+|core\s+)?skills\b",
        r"^technologies\b",
        r"^competencies\b",
        r"^areas\s+of\s+expertise\b"
    ],
    "projects": [
        r"^(key\s+|selected\s+)?projects\b",
        r"^portfolio\b",
        r"^personal\s+projects\b"
    ],
    "certifications": [
        r"^certifications?(\s+and\s+licenses?)?\b",
        r"^licenses?(\s+and\s+certifications?)?\b",
        r"^credentials\b"
    ]
}

MONTH_MAP = {
    "jan": 1, "january": 1,
    "feb": 2, "february": 2,
    "mar": 3, "march": 3,
    "apr": 4, "april": 4,
    "may": 5,
    "jun": 6, "june": 6,
    "jul": 7, "july": 7,
    "aug": 8, "august": 8,
    "sep": 9, "sept": 9, "september": 9,
    "oct": 10, "october": 10,
    "nov": 11, "november": 11,
    "dec": 12, "december": 12
}

def parse_month_name(m_str: str) -> Optional[int]:
    clean = m_str.strip().lower()
    return MONTH_MAP.get(clean)

def parse_date_range_string(text: str) -> Optional[ParsedDateRange]:
    """
    Parses dates like 'Jan 2021 - Present', '2016 - 2018', '03/2019 - 05/2022', 'May 2020 - Current'
    """
    pattern = r'(?P<start>(?:[A-Za-z]{3,9}\.?\s+)?(?:\d{1,2}/)?\d{4})\s*(?:-|–|—|to)\s*(?P<end>present|current|now|(?:[A-Za-z]{3,9}\.?\s+)?(?:\d{1,2}/)?\d{4})'
    match = re.search(pattern, text, re.IGNORECASE)
    if not match:
        # Check single year range like 2018 - 2021
        simple_pat = r'\b(20\d\d|19\d\d)\s*(?:-|–|—|to)\s*(present|current|now|20\d\d|19\d\d)\b'
        match = re.search(simple_pat, text, re.IGNORECASE)
        if not match:
            return None
        s_val, e_val = match.group(1), match.group(2)
        s_yr = int(s_val)
        is_cur = e_val.lower() in ("present", "current", "now")
        e_yr = 2026 if is_cur else int(e_val)
        return ParsedDateRange(raw_text=match.group(0), start_year=s_yr, start_month=1, end_year=e_yr, end_month=12, is_current=is_cur)

    s_str = match.group('start').strip()
    e_str = match.group('end').strip()
    
    # parse start
    s_yr, s_mo = None, 1
    # Check MM/YYYY
    m_slash = re.search(r'(\d{1,2})/(\d{4})', s_str)
    if m_slash:
        s_mo = int(m_slash.group(1))
        s_yr = int(m_slash.group(2))
    else:
        m_word = re.search(r'([A-Za-z]{3,9})\.?\s+(\d{4})', s_str)
        if m_word:
            s_mo = parse_month_name(m_word.group(1)) or 1
            s_yr = int(m_word.group(2))
        else:
            m_yr = re.search(r'\b(20\d\d|19\d\d)\b', s_str)
            if m_yr:
                s_yr = int(m_yr.group(1))

    is_current = e_str.lower() in ("present", "current", "now")
    e_yr, e_mo = (2026, 12) if is_current else (None, 12)
    
    if not is_current:
        m_slash_e = re.search(r'(\d{1,2})/(\d{4})', e_str)
        if m_slash_e:
            e_mo = int(m_slash_e.group(1))
            e_yr = int(m_slash_e.group(2))
        else:
            m_word_e = re.search(r'([A-Za-z]{3,9})\.?\s+(\d{4})', e_str)
            if m_word_e:
                e_mo = parse_month_name(m_word_e.group(1)) or 12
                e_yr = int(m_word_e.group(2))
            else:
                m_yr_e = re.search(r'\b(20\d\d|19\d\d)\b', e_str)
                if m_yr_e:
                    e_yr = int(m_yr_e.group(1))

    if s_yr:
        return ParsedDateRange(
            raw_text=match.group(0),
            start_year=s_yr,
            start_month=s_mo,
            end_year=e_yr,
            end_month=e_mo,
            is_current=is_current
        )
    return None

class ResumeParser:
    """Parses PDF, DOCX, or text files into a structured ParsedResume object."""
    
    @staticmethod
    def extract_bullets_and_dates(text: str) -> Tuple[List[str], List[ParsedDateRange]]:
        bullets = []
        dates = []
        lines = text.split("\n")
        
        for line in lines:
            line_str = line.strip()
            if not line_str:
                continue
            
            # Check date range
            dr = parse_date_range_string(line_str)
            if dr:
                dates.append(dr)
                
            # Check bullet line (starts with -, *, •, ‣, ⁃, or numbered bullet)
            bullet_match = re.match(r'^[\u2022\u2023\u25E6\u2043\u2219\*\-\–\—\>]\s*(.+)$', line_str)
            if bullet_match:
                bullets.append(bullet_match.group(1).strip())
            elif re.match(r'^\d+[\.\)]\s+(.+)$', line_str):
                m = re.match(r'^\d+[\.\)]\s+(.+)$', line_str)
                bullets.append(m.group(1).strip())
            elif len(line_str) > 30 and (line_str.endswith('.') or line_str.endswith(';') or ';' in line_str):
                # Potential bullet formatted without symbol
                bullets.append(line_str)
                
        return bullets, dates

    @staticmethod
    def detect_sections(text: str) -> Dict[str, str]:
        lines = [line.strip() for line in text.split("\n") if line.strip()]
        sections: Dict[str, List[str]] = {}
        current_section = "header"
        sections[current_section] = []
        
        for line in lines:
            normalized = line.lower().strip(":").strip()
            matched_sec = None
            for sec_name, patterns in SECTION_PATTERNS.items():
                for pat in patterns:
                    if re.match(pat, normalized, re.IGNORECASE) and len(normalized.split()) <= 4:
                        matched_sec = sec_name
                        break
                if matched_sec:
                    break
            
            if matched_sec:
                current_section = matched_sec
                if current_section not in sections:
                    sections[current_section] = []
            else:
                sections[current_section].append(line)
                
        return {k: "\n".join(v) for k, v in sections.items() if v}

    @classmethod
    def parse_pdf(cls, file_bytes: bytes) -> ParsedResume:
        raw_text_parts = []
        pages_count = 0
        has_tables = False
        has_images = False
        is_multi_column = False
        total_word_count = 0
        
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            pages_count = len(pdf.pages)
            for page in pdf.pages:
                text = page.extract_text() or ""
                raw_text_parts.append(text)
                words = page.extract_words()
                total_word_count += len(words)
                
                # Check tables
                tables = page.find_tables()
                if tables and len(tables) > 0:
                    has_tables = True
                    
                # Check images
                if hasattr(page, "images") and len(page.images) > 0:
                    has_images = True
                    
                # Multi-column detection: Check distribution of x0 positions of words
                if len(words) > 30:
                    page_width = float(page.width)
                    mid_x = page_width / 2.0
                    left_words = [w for w in words if float(w['x0']) < mid_x * 0.85]
                    right_words = [w for w in words if float(w['x0']) > mid_x * 0.75]
                    # If both left and right have substantial text simultaneously on same vertical band
                    if len(left_words) > 20 and len(right_words) > 20:
                        # Check vertical overlap
                        left_ys = set(int(w['top'] // 20) for w in left_words)
                        right_ys = set(int(w['top'] // 20) for w in right_words)
                        overlap = left_ys.intersection(right_ys)
                        if len(overlap) > 5:
                            is_multi_column = True

        full_text = "\n".join(raw_text_parts)
        # Scanned check: low word count (less than 40 words total across all pages) but images present
        is_scanned = (total_word_count < 40 and has_images)
        
        bullets, dates = cls.extract_bullets_and_dates(full_text)
        sections = cls.detect_sections(full_text)
        
        emails = re.findall(r'[\w\.-]+@[\w\.-]+\.\w+', full_text)
        phones = re.findall(r'(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}', full_text)
        links = re.findall(r'https?://[^\s]+|linkedin\.com/[^\s]+|github\.com/[^\s]+', full_text)
        
        return ParsedResume(
            raw_text=full_text,
            pages=pages_count,
            word_count=total_word_count,
            sections=sections,
            bullet_points=bullets,
            has_tables=has_tables,
            has_images=has_images,
            is_scanned=is_scanned,
            is_multi_column=is_multi_column,
            emails=list(set(emails)),
            phones=list(set(phones)),
            links=list(set(links)),
            date_ranges=dates
        )

    @classmethod
    def parse_docx(cls, file_bytes: bytes) -> ParsedResume:
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs_text = [p.text for p in doc.paragraphs if p.text.strip()]
        
        has_tables = len(doc.tables) > 0
        table_text = []
        for table in doc.tables:
            for row in table.rows:
                row_str = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_str:
                    table_text.append(row_str)
                    
        full_text = "\n".join(paragraphs_text + table_text)
        words = full_text.split()
        word_count = len(words)
        
        # Estimate pages based on word count (approx 450 words per page)
        pages = max(1, (word_count + 400) // 450)
        
        bullets, dates = cls.extract_bullets_and_dates(full_text)
        sections = cls.detect_sections(full_text)
        
        emails = re.findall(r'[\w\.-]+@[\w\.-]+\.\w+', full_text)
        phones = re.findall(r'(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{4}', full_text)
        links = re.findall(r'https?://[^\s]+|linkedin\.com/[^\s]+|github\.com/[^\s]+', full_text)
        
        return ParsedResume(
            raw_text=full_text,
            pages=pages,
            word_count=word_count,
            sections=sections,
            bullet_points=bullets,
            has_tables=has_tables,
            has_images=False,
            is_scanned=False,
            is_multi_column=False,
            emails=list(set(emails)),
            phones=list(set(phones)),
            links=list(set(links)),
            date_ranges=dates
        )

    @classmethod
    def parse_txt(cls, text: str) -> ParsedResume:
        words = text.split()
        word_count = len(words)
        pages = max(1, (word_count + 400) // 450)
        bullets, dates = cls.extract_bullets_and_dates(text)
        sections = cls.detect_sections(text)
        
        emails = re.findall(r'[\w\.-]+@[\w\.-]+\.\w+', text)
        phones = re.findall(r'(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{4}', text)
        links = re.findall(r'https?://[^\s]+|linkedin\.com/[^\s]+|github\.com/[^\s]+', text)
        
        return ParsedResume(
            raw_text=text,
            pages=pages,
            word_count=word_count,
            sections=sections,
            bullet_points=bullets,
            has_tables=False,
            has_images=False,
            is_scanned=False,
            is_multi_column=False,
            emails=list(set(emails)),
            phones=list(set(phones)),
            links=list(set(links)),
            date_ranges=dates
        )

    @classmethod
    def from_resume_data(cls, data: Any) -> ParsedResume:
        """Converts structured ResumeData from Builder into ParsedResume for consistent analysis."""
        lines = []
        if hasattr(data, "contact") and data.contact:
            c = data.contact
            lines.append(f"{c.full_name} | {c.job_title} | {c.email} | {c.phone} | {c.location} | {c.linkedin} | {c.github} | {c.website}")
            
        sections = {}
        if hasattr(data, "summary") and data.summary:
            sections["summary"] = data.summary
            lines.append("SUMMARY\n" + data.summary)
            
        exp_lines = []
        bullets = []
        date_ranges = []
        if hasattr(data, "experience") and data.experience:
            for exp in data.experience:
                date_str = f"{exp.start_date} - {'Present' if exp.current else exp.end_date}"
                exp_lines.append(f"{exp.title} at {exp.company} ({date_str})")
                dr = parse_date_range_string(date_str)
                if dr:
                    date_ranges.append(dr)
                for b in exp.bullets:
                    if b.strip():
                        exp_lines.append(f"• {b.strip()}")
                        bullets.append(b.strip())
            sections["experience"] = "\n".join(exp_lines)
            lines.append("EXPERIENCE\n" + sections["experience"])
            
        proj_lines = []
        if hasattr(data, "projects") and data.projects:
            for p in data.projects:
                proj_lines.append(f"{p.title}: {p.description} (Tech: {', '.join(p.technologies)})")
            sections["projects"] = "\n".join(proj_lines)
            lines.append("PROJECTS\n" + sections["projects"])
            
        edu_lines = []
        if hasattr(data, "education") and data.education:
            for ed in data.education:
                edu_lines.append(f"{ed.degree} in {ed.field_of_study}, {ed.institution} ({ed.start_date} - {ed.end_date})")
                dr = parse_date_range_string(f"{ed.start_date} - {ed.end_date}")
                if dr:
                    date_ranges.append(dr)
            sections["education"] = "\n".join(edu_lines)
            lines.append("EDUCATION\n" + sections["education"])
            
        skills_lines = []
        if hasattr(data, "skills") and data.skills:
            skills_lines.append(", ".join([s.name for s in data.skills if s.name]))
            sections["skills"] = "\n".join(skills_lines)
            lines.append("SKILLS\n" + sections["skills"])
            
        cert_lines = []
        if hasattr(data, "certifications") and data.certifications:
            for cert in data.certifications:
                cert_lines.append(f"{cert.name} by {cert.issuer} ({cert.issue_date})")
            sections["certifications"] = "\n".join(cert_lines)
            lines.append("CERTIFICATIONS\n" + sections["certifications"])
            
        full_text = "\n\n".join(lines)
        words = full_text.split()
        word_count = len(words)
        pages = max(1, (word_count + 400) // 450)
        
        emails = [data.contact.email] if hasattr(data, "contact") and data.contact and data.contact.email else []
        phones = [data.contact.phone] if hasattr(data, "contact") and data.contact and data.contact.phone else []
        links = []
        if hasattr(data, "contact") and data.contact:
            if data.contact.linkedin: links.append(data.contact.linkedin)
            if data.contact.github: links.append(data.contact.github)
            if data.contact.website: links.append(data.contact.website)
            
        return ParsedResume(
            raw_text=full_text,
            pages=pages,
            word_count=word_count,
            sections=sections,
            bullet_points=bullets,
            has_tables=False,
            has_images=False,
            is_scanned=False,
            is_multi_column=False,
            emails=emails,
            phones=phones,
            links=links,
            date_ranges=date_ranges
        )
