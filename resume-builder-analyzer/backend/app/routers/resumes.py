from fastapi import APIRouter, Depends, HTTPException, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from typing import List, Optional
import json
import io
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from app.database import get_db
from app.models import User, Resume
from app.schemas import ResumeCreate, ResumeUpdate, ResumeOut, ResumeData
from app.services.auth import get_current_user, get_optional_current_user

router = APIRouter(prefix="/resumes", tags=["Resume Builder"])

def serialize_resume(resume: Resume) -> ResumeOut:
    data_dict = json.loads(resume.data_json)
    return ResumeOut(
        id=resume.id,
        user_id=resume.user_id,
        title=resume.title,
        template_id=resume.template_id,
        accent_color=resume.accent_color,
        font_family=resume.font_family,
        data=ResumeData.model_validate(data_dict),
        created_at=resume.created_at,
        updated_at=resume.updated_at
    )

@router.get("", response_model=List[ResumeOut])
def list_resumes(current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    resumes = db.query(Resume).filter(Resume.user_id == current_user.id).order_by(Resume.updated_at.desc()).all()
    return [serialize_resume(r) for r in resumes]

@router.post("", response_model=ResumeOut, status_code=status.HTTP_201_CREATED)
def create_resume(resume_in: ResumeCreate, current_user: Optional[User] = Depends(get_optional_current_user), db: Session = Depends(get_db)):
    user_id = current_user.id if current_user else None
    new_resume = Resume(
        user_id=user_id,
        title=resume_in.title,
        template_id=resume_in.template_id,
        accent_color=resume_in.accent_color,
        font_family=resume_in.font_family,
        data_json=json.dumps(resume_in.data.model_dump())
    )
    db.add(new_resume)
    db.commit()
    db.refresh(new_resume)
    return serialize_resume(new_resume)

@router.get("/{resume_id}", response_model=ResumeOut)
def get_resume(resume_id: int, current_user: Optional[User] = Depends(get_optional_current_user), db: Session = Depends(get_db)):
    resume = db.query(Resume).filter(Resume.id == resume_id).first()
    if not resume:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Resume not found")
    if resume.user_id and (not current_user or resume.user_id != current_user.id):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Access denied")
    return serialize_resume(resume)

@router.put("/{resume_id}", response_model=ResumeOut)
def update_resume(resume_id: int, resume_in: ResumeUpdate, current_user: Optional[User] = Depends(get_optional_current_user), db: Session = Depends(get_db)):
    resume = db.query(Resume).filter(Resume.id == resume_id).first()
    if not resume:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Resume not found")
    if resume.user_id and (not current_user or resume.user_id != current_user.id):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Access denied")
        
    if resume_in.title is not None:
        resume.title = resume_in.title
    if resume_in.template_id is not None:
        resume.template_id = resume_in.template_id
    if resume_in.accent_color is not None:
        resume.accent_color = resume_in.accent_color
    if resume_in.font_family is not None:
        resume.font_family = resume_in.font_family
    if resume_in.data is not None:
        resume.data_json = json.dumps(resume_in.data.model_dump())
        
    db.commit()
    db.refresh(resume)
    return serialize_resume(resume)

@router.post("/{resume_id}/duplicate", response_model=ResumeOut, status_code=status.HTTP_201_CREATED)
def duplicate_resume(resume_id: int, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    original = db.query(Resume).filter(Resume.id == resume_id, Resume.user_id == current_user.id).first()
    if not original:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Resume not found")
        
    duplicated = Resume(
        user_id=current_user.id,
        title=f"{original.title} (Copy)",
        template_id=original.template_id,
        accent_color=original.accent_color,
        font_family=original.font_family,
        data_json=original.data_json
    )
    db.add(duplicated)
    db.commit()
    db.refresh(duplicated)
    return serialize_resume(duplicated)

@router.delete("/{resume_id}", status_code=status.HTTP_204_NO_CONTENT)
def delete_resume(resume_id: int, current_user: User = Depends(get_current_user), db: Session = Depends(get_db)):
    resume = db.query(Resume).filter(Resume.id == resume_id, Resume.user_id == current_user.id).first()
    if not resume:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Resume not found")
    db.delete(resume)
    db.commit()
    return None

@router.post("/export/docx")
def export_docx(data: ResumeData):
    """Generates an ATS-compliant, well-styled Microsoft Word (.docx) resume file."""
    doc = docx.Document()
    
    # Page Margins 0.7 inch
    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.7)
        section.right_margin = Inches(0.7)

    # Header Name & Title
    c = data.contact
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_name = p_name.add_run(c.full_name or "Your Name")
    run_name.font.size = Pt(20)
    run_name.font.bold = True
    run_name.font.name = "Calibri"
    
    if c.job_title:
        p_title = doc.add_paragraph()
        p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_title = p_title.add_run(c.job_title)
        run_title.font.size = Pt(13)
        run_title.font.italic = True
        run_title.font.color.rgb = RGBColor(80, 80, 80)
        
    # Contact Details line
    contact_parts = [p for p in [c.email, c.phone, c.location, c.linkedin, c.github, c.website] if p]
    if contact_parts:
        p_contact = doc.add_paragraph()
        p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run_contact = p_contact.add_run(" | ".join(contact_parts))
        run_contact.font.size = Pt(9.5)
        run_contact.font.color.rgb = RGBColor(100, 100, 100)

    def add_section_header(title: str):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(title.upper())
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(30, 41, 59)

    # Summary
    if data.summary:
        add_section_header("Professional Summary")
        p_sum = doc.add_paragraph(data.summary)
        p_sum.paragraph_format.space_after = Pt(6)
        p_sum.style.font.size = Pt(10)

    # Experience
    if data.experience:
        add_section_header("Work Experience")
        for exp in data.experience:
            p_role = doc.add_paragraph()
            p_role.paragraph_format.space_before = Pt(4)
            p_role.paragraph_format.space_after = Pt(1)
            
            run_role = p_role.add_run(exp.title)
            run_role.bold = True
            run_role.font.size = Pt(10.5)
            
            run_comp = p_role.add_run(f" | {exp.company}")
            run_comp.font.size = Pt(10)
            
            date_str = f" ({exp.start_date} – {'Present' if exp.current else exp.end_date})"
            run_date = p_role.add_run(date_str)
            run_date.italic = True
            run_date.font.size = Pt(9.5)
            run_date.font.color.rgb = RGBColor(100, 100, 100)
            
            for b in exp.bullets:
                if b.strip():
                    p_b = doc.add_paragraph(style='List Bullet')
                    p_b.paragraph_format.space_after = Pt(1)
                    p_b.paragraph_format.space_before = Pt(0)
                    r = p_b.add_run(b.strip())
                    r.font.size = Pt(9.5)

    # Projects
    if data.projects:
        add_section_header("Key Projects")
        for proj in data.projects:
            p_proj = doc.add_paragraph()
            p_proj.paragraph_format.space_before = Pt(4)
            p_proj.paragraph_format.space_after = Pt(1)
            run_pname = p_proj.add_run(proj.title)
            run_pname.bold = True
            run_pname.font.size = Pt(10)
            
            if proj.technologies:
                run_tech = p_proj.add_run(f" [{', '.join(proj.technologies)}]")
                run_tech.font.size = Pt(9)
                run_tech.font.color.rgb = RGBColor(100, 100, 100)
                
            p_desc = doc.add_paragraph(style='List Bullet')
            p_desc.paragraph_format.space_after = Pt(1)
            r = p_desc.add_run(proj.description)
            r.font.size = Pt(9.5)

    # Education
    if data.education:
        add_section_header("Education")
        for edu in data.education:
            p_edu = doc.add_paragraph()
            p_edu.paragraph_format.space_before = Pt(3)
            p_edu.paragraph_format.space_after = Pt(1)
            r_deg = p_edu.add_run(f"{edu.degree} in {edu.field_of_study}")
            r_deg.bold = True
            r_deg.font.size = Pt(10)
            
            r_inst = p_edu.add_run(f" – {edu.institution}")
            r_inst.font.size = Pt(9.5)
            
            if edu.start_date or edu.end_date:
                r_dates = p_edu.add_run(f" ({edu.start_date} – {edu.end_date})")
                r_dates.font.size = Pt(9)
                r_dates.font.color.rgb = RGBColor(100, 100, 100)

    # Skills
    if data.skills:
        add_section_header("Skills & Competencies")
        p_sk = doc.add_paragraph()
        p_sk.paragraph_format.space_after = Pt(4)
        skill_names = [s.name for s in data.skills if s.name]
        r_sk = p_sk.add_run(" • ".join(skill_names))
        r_sk.font.size = Pt(9.5)

    # Certifications
    if data.certifications:
        add_section_header("Certifications")
        for cert in data.certifications:
            p_c = doc.add_paragraph(style='List Bullet')
            p_c.paragraph_format.space_after = Pt(1)
            r_c = p_c.add_run(f"{cert.name} – {cert.issuer} ({cert.issue_date})")
            r_c.font.size = Pt(9.5)

    stream = io.BytesIO()
    doc.save(stream)
    stream.seek(0)
    
    filename = f"{data.contact.full_name.replace(' ', '_') or 'Resume'}.docx"
    return StreamingResponse(
        stream,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )
